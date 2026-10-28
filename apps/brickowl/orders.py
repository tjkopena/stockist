
if __name__ != '__main__':
    raise Error("This is a script, not intended to be imported")

from datetime import datetime, timedelta
from collections import defaultdict
import argparse
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from stockist import brickowl

import logberry
import tjawn


CURRENCY = '"$"#,##0.00'
DATETIME = 'yyyy-mm-dd hh:mm'

SUMMARY_COLUMNS = [
    ('BOID', 14, None),
    ('Natural ID', 14, None),
    ('Name', 60, None),
    ('Color', 20, None),
    ('Quantity', 10, None),
    ('Min Adj. Price', 14, CURRENCY),
    ('Max Adj. Price', 14, CURRENCY),
    ('Median Adj. Price', 14, CURRENCY),
    ('Mean Adj. Price', 14, CURRENCY),
    ('Adj. Total', 14, CURRENCY),
    ('Min Price', 14, CURRENCY),
    ('Max Price', 14, CURRENCY),
    ('Median Price', 14, CURRENCY),
    ('Mean Price', 14, CURRENCY),
    ('Total', 14, CURRENCY),
]

ORDER_COLUMNS = [
    ('BOID', 24, None),
    ('Natural ID', 14, None),
    ('Name', 60, None),
    ('Color', 20, None),
    ('Condition', 16, None),
    ('Quantity', 10, None),
    ('Adj. Price', 14, CURRENCY),
    ('Adj. Subtotal', 14, CURRENCY),
    ('Price', 14, CURRENCY),
    ('Subtotal', 14, CURRENCY),
]

ORDER_SUMMARY_FORMATS = [None, None, DATETIME, None, None,
                         CURRENCY, CURRENCY, CURRENCY, CURRENCY, CURRENCY]

SHEET_NAME_MAX = 31


def parse_date(s):
    for fmt in ('%Y%m%d', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise argparse.ArgumentTypeError(f"Invalid date '{s}', must be YYYYMMDD or YYYY-MM-DD")


parser = argparse.ArgumentParser(prog='python -m apps.brickowl.orders',
                                 description='Summarize Brick Owl purchase orders into an XLSX workbook')
parser.add_argument('--key', help='Brick Owl API key (default from .conf/brickowl.jawn)')
parser.add_argument('--orders', help='Comma separated list of order IDs')
parser.add_argument('--start', type=parse_date, help='Earliest order date, YYYYMMDD or YYYY-MM-DD')
parser.add_argument('--end', type=parse_date, help='Latest order date (inclusive), YYYYMMDD or YYYY-MM-DD')
parser.add_argument('--output', help='Output XLSX filename (default data/YYYYMMDD.xlsx)')
args = parser.parse_args()

if args.orders and (args.start or args.end):
    parser.error("--orders cannot be combined with --start or --end")

today = datetime.now().strftime('%Y%m%d')
output = args.output or f"data/{today}.xlsx"
fn_raw = f"data/brickowl-orders-{today}.jawn"

logberry.start()

settings = brickowl.Settings(api_key=args.key) if args.key else brickowl.load_configuration()
api = brickowl.client(settings)


#-- Download --------------------------------------------------------------

if args.orders:
    order_ids = [o.strip() for o in args.orders.split(',') if o.strip()]
else:
    start_ts = int(args.start.timestamp()) if args.start else None
    end_ts = int((args.end + timedelta(days=1)).timestamp()) if args.end else None
    order_ids = [o['order_id'] for o in api.list_orders(order_time=start_ts)
                 if end_ts is None or int(o['order_date']) < end_ts]

t = logberry.task("Download orders", orders=len(order_ids))
raw = []
for oid in order_ids:
    view = api.view_order(oid)
    if brickowl.is_cancelled(view):
        t.info("Skipping cancelled order", order_id=oid)
        continue
    raw.append({'view': view, 'items': api.order_items(oid)})
t.end_success(orders=len(raw))

all_items = [i for o in raw for i in o['items']]
bl_ids = (api.id_list('Minifigure', 'bl_item_no')
          if any(i['type'] == 'Minifigure' for i in all_items) else {})
design_ids = (api.id_list('Part', 'design_id')
              if any(brickowl.needs_design_ids(i) for i in all_items) else {})

t = logberry.task("Write raw orders", file=fn_raw)
with open(fn_raw, "wt") as f:
    f.write(tjawn.dumps(raw))
t.end_success()


#-- Process ---------------------------------------------------------------

orders = [brickowl.as_order(o['view'], o['items'], brickowl.order_rate(o['view']),
                            design_ids, bl_ids)
          for o in raw]
orders.sort(key=lambda o: o.placed, reverse=True)

grouped = defaultdict(list)
for o in orders:
    for i in o.items:
        grouped[i.key].append(i)

summary = []
for items in grouped.values():
    first = items[0]
    adjusted = brickowl.weighted_stats([(i.adjusted_price, i.quantity) for i in items])
    plain = brickowl.weighted_stats([(i.price, i.quantity) for i in items])
    summary.append([first.boid, first.natural_id, first.name, first.color,
                    sum(i.quantity for i in items), *adjusted, *plain])
summary.sort(key=lambda r: r[9], reverse=True)


#-- Sheet names -----------------------------------------------------------

def sanitize(name):
    name = re.sub(r"[\[\]:*?/\\']", '', name)
    return re.sub(r'\s+', ' ', name).strip()

same_day = defaultdict(list)
for o in orders:
    same_day[(sanitize(o.store_name).lower(), o.placed.date())].append(o)

sheet_names = {}
for group in same_day.values():
    for x, o in enumerate(group):
        suffix = f" ({len(group) - x})" if len(group) > 1 else ''
        stamp = f" {o.placed:%Y%m%d}{suffix}"
        store = sanitize(o.store_name)[:SHEET_NAME_MAX - len(stamp)].strip()
        sheet_names[o.order_id] = f"{store}{stamp}"


#-- Workbook --------------------------------------------------------------

bold = Font(bold=True)

def header(ws, row, columns):
    for c, (label, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = bold
        ws.column_dimensions[cell.column_letter].width = width

def rows(ws, start, data, formats):
    for r, values in enumerate(data, start=start):
        for c, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if fmt:
                cell.number_format = fmt

t = logberry.task("Write workbook", file=output)
wb = Workbook()

ws = wb.active
ws.title = 'Summary'
header(ws, 1, SUMMARY_COLUMNS)
rows(ws, 2, summary, [f for _, _, f in SUMMARY_COLUMNS])
ws.freeze_panes = 'A2'

for o in orders:
    ws = wb.create_sheet(sheet_names[o.order_id])
    rows(ws, 1, [[o.store_name, int(o.order_id), o.placed, o.lots, o.quantity,
                  o.subtotal, o.shipping, o.taxes, o.discounts, o.total]],
         ORDER_SUMMARY_FORMATS)
    for cell in ws[1]:
        cell.font = bold
    ws.cell(row=1, column=8).comment = Comment(
        "Taxes are informational: VAT/sales tax is already included in item prices and shipping.",
        "stockist")

    header(ws, 2, ORDER_COLUMNS)
    rows(ws, 3, [[i.boid, i.natural_id, i.name, i.color, i.condition, i.quantity,
                  i.adjusted_price, i.adjusted_subtotal, i.price, i.subtotal]
                 for i in o.items],
         [f for _, _, f in ORDER_COLUMNS])
    ws.freeze_panes = 'A3'

wb.save(output)
t.end_success(orders=len(orders), items=len(summary))

logberry.stop()
